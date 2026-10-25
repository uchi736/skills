# -*- coding: utf-8 -*-
"""
make_ideal_pairs — 直した理想見本キット(.pptx)を assets/ideal/ の before/after PNG ペアにする。

    python make_ideal_pairs.py <kit.pptx> [out_dir]   (既定: <このスキル>/assets/ideal)

- キットの各枚には make_ideal_kit.py が付けた図形 "IDEAL_MARKER|BEFORE|<カテゴリ>" /
  "IDEAL_MARKER|AFTER|<カテゴリ>" がある前提。消えていた場合は並び順(前後交互)で補う
- 目印の図形は PNG に写らないよう、一時フォルダのコピーから外してからレンダリングする
- 出力: NN_<カテゴリ>_before.png / NN_<カテゴリ>_after.png / index.md(型ごとの一覧)
- 既存の見本は消さない。番号は既存の最大の続きから振り、index.md には行を追記する
"""
import os
import re
import shutil
import sys
import tempfile

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from export_preview import export_pngs, _winpath  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_OUT = os.path.join(HERE, "..", "assets", "ideal")
PNG_RE = re.compile(r"^(\d{2,})_.+_(before|after)\.png$")
HEADER = [
    "# 理想見本(before/after ペア)",
    "",
    "組む前に、作る型に近いペアを Read で見比べてから書く。after が基準。",
    "before は「今のスキルが出しがちな形」なので、差分(文の長さ・図解の密度・アイコン・要旨の語り方)を適用する。",
    "見本と SKILL.md の鉄則(文字 12pt 以上・▍1 文リード・kpi タイル不使用など)が食い違うときは SKILL.md が優先。",
    "",
    "| No. | カテゴリ | 元スライド | before | after |",
    "|---|---|---|---|---|",
]
FOOT = ("見本の補足: 公式 IR 資料は `ir-fy2026/README.md`(要旨 p03、表の代替 p06・p13・p18、結び p16)、"
        "コンサル実物は `abeam/README.md`(パネル+番号凡例+構成図、左右の問いのグラフ)。")


def read_entries(kit, tmpdir):
    """各枚の (role, category, title) を読み、目印を外したコピーを tmpdir に保存して返す。"""
    prs = Presentation(kit)
    entries = []
    prev_role, prev_cat = "AFTER", ""
    for s in prs.slides:
        role = cat = None
        for sh in list(s.shapes):
            if sh.name.startswith("IDEAL_MARKER|"):
                _, role, cat = sh.name.split("|", 2)
                sh._element.getparent().remove(sh._element)
        if role is None:                       # 目印が消えていたら交互に補う
            role = "AFTER" if prev_role == "BEFORE" else "BEFORE"
            cat = prev_cat
        title = ""
        if s.shapes.title is not None:
            title = s.shapes.title.text_frame.text.strip()
        entries.append((role, cat, title))
        prev_role, prev_cat = role, cat
    stripped = os.path.join(tmpdir, "kit_nomarker.pptx")
    prs.save(stripped)
    return entries, stripped


def _load_index(idx_path):
    """既存 index.md を (表までの行, 補足行) に分ける。空の案内行は捨てる。"""
    if not os.path.exists(idx_path):
        return list(HEADER), []
    with open(idx_path, encoding="utf-8") as f:
        old = f.read().rstrip("\n").split("\n")
    body = [ln for ln in old if "(まだ無い" not in ln]
    last_row = max((i for i, ln in enumerate(body) if ln.startswith("|")), default=None)
    if last_row is None:
        return list(HEADER), []
    tail = [ln for ln in body[last_row + 1:] if ln.strip()
            and not ln.startswith("公式IR資料のページ一覧")]
    return body[:last_row + 1], tail


def main(kit, out_dir):
    kit = _winpath(kit)
    out_dir = _winpath(out_dir)
    if not os.path.isfile(kit):
        raise SystemExit(f"キットがありません: {kit}")
    tmpdir = tempfile.mkdtemp(prefix="ideal_pairs_")
    try:
        entries, stripped = read_entries(kit, tmpdir)
        pngs = export_pngs(stripped, os.path.join(tmpdir, "png"))
        if len(pngs) != len(entries):
            raise SystemExit(f"PNG {len(pngs)} 枚とスライド {len(entries)} 枚が一致しない")
        pairs = []
        k = 0
        while k < len(entries) - 1:
            b, a = entries[k], entries[k + 1]
            if b[0] == "BEFORE" and a[0] == "AFTER":
                pairs.append((k, k + 1))
                k += 2
            else:
                k += 1
        if not pairs:
            raise SystemExit("BEFORE/AFTER の対が見つからない(目印の図形名を確認)")
        os.makedirs(out_dir, exist_ok=True)
        start = max([int(m.group(1)) for f in os.listdir(out_dir)
                     for m in [PNG_RE.match(f)] if m] + [0])
        idx_path = os.path.join(out_dir, "index.md")
        lines, tail = _load_index(idx_path)
        for n, (bi, ai) in enumerate(pairs, start + 1):
            _, cat, title = entries[bi]
            safe = re.sub(r'[\\/:*?"<>| ]', "_", cat) or "misc"
            bn, an = f"{n:02d}_{safe}_before.png", f"{n:02d}_{safe}_after.png"
            shutil.copy(pngs[bi], os.path.join(out_dir, bn))
            shutil.copy(pngs[ai], os.path.join(out_dir, an))
            lines.append(f"| {n:02d} | {cat} | {title} | {bn} | {an} |")
        if FOOT not in tail:
            tail.append(FOOT)
        with open(idx_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines + [""] + tail) + "\n")
        print(f"pairs: {len(pairs)}(No.{start + 1:02d}〜{start + len(pairs):02d})-> {out_dir}")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if len(sys.argv) < 2 or sys.argv[1] in ("-h", "--help"):
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT)
