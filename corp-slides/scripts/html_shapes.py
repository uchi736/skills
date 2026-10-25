# -*- coding: utf-8 -*-
"""
html_shapes — HTML/CSS で組んだ図解を、PowerPoint で編集できるネイティブ図形に変換する。

    from html_shapes import html_diagram, place_html
    html_diagram(d, "タイトル", BODY_HTML, css=CSS, lead="主張の 1 文", source="出典", keep="figs/p6.html")

仕組み: ヘッドレス Edge / Chrome で HTML を本文領域と同じ大きさ(inch)で描画し、要素の位置・色・枠・文字の行を
測って JSON にし、python-pptx の図形(矩形・楕円・テキストボックス・コネクタ・フリーフォーム・線画アイコン)で
置き直す。画像にはしないので、箱・文字・矢印は全部 PowerPoint で直せる。

■ HTML の書き方(図のキャンバス = 本文領域。#fig の中に書く。幅・高さは呼び出し側が決める)
  - 文字は pt で書く(14pt 本文 / 12pt 補足が下限)。行の折り返し位置はブラウザの結果をそのまま再現する
  - 背景色か枠線のある要素 → 矩形(border-radius ≥ 50% は楕円)。片側だけの枠 → その辺の線
  - border-style: dashed → 破線
  - 文字を持つブロック要素 → テキストボックス(太字・色・大きさは run ごとに引き継ぐ)。::before の文字は拾わない
  - <i class="ico" data-icon="db"></i>(幅・高さ指定)→ 線画アイコン 13 種(corp_pptx.icon)
  - <x-wire from="#a" to="#b" route="h|v|hv|vh|auto" label="渡すもの" dash head="end|start|both|none"
            color="#3377AE" width="1.75" gap="4"> → 矢印(2 点はコネクタ、折れ線はフリーフォーム)
    from/to は "#id" か "#id:r" / "#id:b:0.3"(辺 l r t b と辺上の位置 0〜1)。points="x,y x,y …"(px)も可
  - 使えるクラス(既定 CSS): .box .hl .plan .head .zone .badge .lbl .sub .cap(下記 BASE_CSS)
■ 色は corp_pptx のパレットを CSS 変数で使う(var(--blue) など)。純黒は使わない
"""
import html as _html
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from pptx.util import Inches, Pt, Emu  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pptx.enum.dml import MSO_LINE  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
import corp_pptx as C  # noqa: E402
from html2png import _candidates  # noqa: E402

PX = 96.0          # CSS px / inch
PT_PER_PX = 0.75   # 1px = 0.75pt

BASE_CSS = """
:root{--navy:#01357A;--blue:#004EA2;--bluetx:#00549A;--steel:#3377AE;--cyan:#00BCEB;--teal:#00BB99;
 --red:#C00000;--bg:#F0F5F8;--ltblue:#D6ECF7;--pale:#E8F2FA;--ink:#262626;--ink2:#4D4D4D;--grey:#A6A6A6;--grid:#E3E7EC}
html,body{margin:0;padding:0;background:#fff}
#fig{position:relative;overflow:visible;box-sizing:border-box;font-family:"Yu Gothic","YuGothic",sans-serif;
 font-size:14pt;line-height:1.4;color:var(--ink);font-weight:400}
#fig *{box-sizing:border-box}
#fig b,#fig strong{font-weight:700}
.box{background:#fff;border:0.75pt solid var(--grey);padding:6pt 9pt}
.hl{background:var(--blue);border-color:var(--blue);color:#fff;font-weight:700}
.plan{border-style:dashed;color:var(--ink2)}
.head{background:var(--blue);color:#fff;font-weight:700;padding:4pt 9pt;text-align:center}
.zone{border:1.25pt dashed var(--steel);background:transparent}
.badge{display:inline-flex;align-items:center;justify-content:center;width:0.34in;height:0.34in;border-radius:50%;
 background:var(--navy);color:#fff;font-size:12pt;font-weight:700;line-height:1}
.lbl{font-size:12pt;color:var(--ink2)}
.sub{font-size:12pt;color:var(--ink2)}
.cap{font-weight:700;color:var(--bluetx)}
.ico{display:block}
x-wire{display:none}
.__wl{position:absolute;white-space:nowrap;font-size:12pt;color:var(--ink2);line-height:1.2}
"""

_MEASURE_JS = r"""
(function(){
const fig=document.getElementById('fig');
const F=fig.getBoundingClientRect();
const out={w:F.width,h:F.height,items:[],wires:[],warn:[]};
function col(c){const m=(c||'').match(/rgba?\(([^)]+)\)/);if(!m)return null;
 const p=m[1].split(',').map(s=>parseFloat(s));if(p.length===4&&p[3]<0.05)return null;
 return '#'+p.slice(0,3).map(v=>Math.round(v).toString(16).padStart(2,'0')).join('').toUpperCase();}
function rel(r){return {x:r.left-F.left,y:r.top-F.top,w:r.width,h:r.height};}
function inl(el){const d=getComputedStyle(el).display;return d==='inline'||d==='contents';}
// ---- wires: 端点を決め、ラベルは実際の DOM 要素にして文字として測る
function sideOf(rc,side,t){t=(t==null?0.5:t);
 if(side==='r')return [rc.x+rc.w,rc.y+rc.h*t];if(side==='l')return [rc.x,rc.y+rc.h*t];
 if(side==='t')return [rc.x+rc.w*t,rc.y];return [rc.x+rc.w*t,rc.y+rc.h];}
function parseEnd(s){const p=s.split(':');const el=fig.querySelector(p[0]);
 if(!el){out.warn.push('wire の要素が無い: '+s);return null;}
 return {rc:rel(el.getBoundingClientRect()),side:p[1]||null,t:(p[2]!=null?parseFloat(p[2]):null)};}
const OUT={r:[1,0],l:[-1,0],t:[0,-1],b:[0,1]};
for(const w of fig.querySelectorAll('x-wire')){
 let pts=null;const gap=parseFloat(w.getAttribute('gap')||'4');
 if(w.getAttribute('points')){pts=w.getAttribute('points').trim().split(/\s+/).map(q=>q.split(',').map(Number));}
 else{
  const A=parseEnd(w.getAttribute('from')),B=parseEnd(w.getAttribute('to'));if(!A||!B)continue;
  const a=A.rc,b=B.rc,acx=a.x+a.w/2,acy=a.y+a.h/2,bcx=b.x+b.w/2,bcy=b.y+b.h/2;
  let route=w.getAttribute('route')||'auto';
  if(route==='auto')route=(b.x>=a.x+a.w-1||a.x>=b.x+b.w-1)?'h':'v';
  let sa=A.side,sb=B.side;
  if(route==='h'){sa=sa||(bcx>acx?'r':'l');sb=sb||(bcx>acx?'l':'r');}
  else if(route==='v'){sa=sa||(bcy>acy?'b':'t');sb=sb||(bcy>acy?'t':'b');}
  else if(route==='hv'){sa=sa||(bcx>acx?'r':'l');sb=sb||(bcy>acy?'t':'b');}
  else if(route==='vh'){sa=sa||(bcy>acy?'b':'t');sb=sb||(bcx>acx?'l':'r');}
  let ta=A.t,tb=B.t;
  // 真っすぐ引ける重なりがあれば、重なりの中央で直線にする
  if(route==='h'&&ta==null&&tb==null){const lo=Math.max(a.y,b.y),hi=Math.min(a.y+a.h,b.y+b.h);
   if(hi-lo>8){const ym=(lo+hi)/2;ta=(ym-a.y)/a.h;tb=(ym-b.y)/b.h;}}
  if(route==='v'&&ta==null&&tb==null){const lo=Math.max(a.x,b.x),hi=Math.min(a.x+a.w,b.x+b.w);
   if(hi-lo>8){const xm=(lo+hi)/2;ta=(xm-a.x)/a.w;tb=(xm-b.x)/b.w;}}
  let P=sideOf(a,sa,ta),Q=sideOf(b,sb,tb);
  P=[P[0]+OUT[sa][0]*gap,P[1]+OUT[sa][1]*gap];Q=[Q[0]+OUT[sb][0]*gap,Q[1]+OUT[sb][1]*gap];
  if(route==='h'){pts=Math.abs(P[1]-Q[1])<1?[P,[Q[0],P[1]]]:[P,[(P[0]+Q[0])/2,P[1]],[(P[0]+Q[0])/2,Q[1]],Q];}
  else if(route==='v'){pts=Math.abs(P[0]-Q[0])<1?[P,[P[0],Q[1]]]:[P,[P[0],(P[1]+Q[1])/2],[Q[0],(P[1]+Q[1])/2],Q];}
  else if(route==='hv'){pts=[P,[Q[0],P[1]],Q];}
  else {pts=[P,[P[0],Q[1]],Q];}
 }
 const rec={pts:pts,color:(w.getAttribute('color')||'#3377AE'),width:parseFloat(w.getAttribute('width')||'1.75'),
  dash:w.hasAttribute('dash'),head:(w.getAttribute('head')||'end')};
 out.wires.push(rec);
 const lab=w.getAttribute('label');
 if(lab){let bi=0,bl=-1;for(let i=0;i<pts.length-1;i++){const L=Math.hypot(pts[i+1][0]-pts[i][0],pts[i+1][1]-pts[i][1]);if(L>bl){bl=L;bi=i;}}
  const p0=pts[bi],p1=pts[bi+1],mx=(p0[0]+p1[0])/2,my=(p0[1]+p1[1])/2;
  const d=document.createElement('div');d.className='__wl';d.textContent=lab;
  if(w.getAttribute('label-class'))d.className+=' '+w.getAttribute('label-class');
  fig.appendChild(d);const lw=d.offsetWidth,lh=d.offsetHeight;const pos=w.getAttribute('label-pos');
  if(Math.abs(p0[1]-p1[1])<1){ // 横の線
   d.style.left=(mx-lw/2)+'px';d.style.top=(pos==='below'?my+3:my-lh-2)+'px';}
  else{d.style.top=(my-lh/2)+'px';d.style.left=(pos==='left'?mx-lw-6:mx+6)+'px';}}
}
// ---- boxes と文字
const rg=document.createRange();
function emitBox(el,cs){
 const r=rel(el.getBoundingClientRect());if(r.w<0.5||r.h<0.5)return;
 const fill=col(cs.backgroundColor);
 const sides=['Top','Right','Bottom','Left'].map(k=>({w:parseFloat(cs['border'+k+'Width'])||0,
   c:col(cs['border'+k+'Color']),s:cs['border'+k+'Style']}));
 const vis=sides.map(sd=>sd.w>0&&sd.c&&sd.s!=='none'&&sd.s!=='hidden');
 const uni=vis.every(v=>v)&&sides.every(sd=>sd.w===sides[0].w&&sd.c===sides[0].c&&sd.s===sides[0].s);
 const rad=parseFloat(cs.borderTopLeftRadius)||0;
 if(!fill&&!vis.some(v=>v))return;
 const rec={type:'rect',x:r.x,y:r.y,w:r.w,h:r.h,fill:fill,radius:rad,ellipse:rad>=Math.min(r.w,r.h)/2-0.5};
 if(uni){rec.line=sides[0].c;rec.lw=sides[0].w;rec.dash=(sides[0].s==='dashed'||sides[0].s==='dotted');}
 out.items.push(rec);
 if(!uni){const E=[[r.x,r.y,r.x+r.w,r.y],[r.x+r.w,r.y,r.x+r.w,r.y+r.h],[r.x,r.y+r.h,r.x+r.w,r.y+r.h],[r.x,r.y,r.x,r.y+r.h]];
  sides.forEach((sd,i)=>{if(!vis[i])return;const e=E[i];const o=sd.w/2;
   const off=[[0,o],[-o,0],[0,-o],[o,0]][i];
   out.items.push({type:'line',x1:e[0]+off[0],y1:e[1]+off[1],x2:e[2]+off[0],y2:e[3]+off[1],color:sd.c,lw:sd.w,
     dash:(sd.s==='dashed'||sd.s==='dotted')});});}
}
function collect(node,acc){for(const n of node.childNodes){
 if(n.nodeType===3)acc.push(n);
 else if(n.nodeType===1){if(n.tagName==='BR'){acc.push('BR');continue;}
  if(n.tagName==='X-WIRE')continue;if(inl(n))collect(n,acc);}}}
function emitText(el,cs){
 const acc=[];collect(el,acc);
 if(!acc.some(n=>n!=='BR'&&n.data.trim()))return;
 const lines=[];let cur=null,forced=false;
 for(const n of acc){
  if(n==='BR'){forced=true;continue;}
  const st=getComputedStyle(n.parentElement);
  const sty={size:parseFloat(st.fontSize)*0.75,bold:(parseInt(st.fontWeight)||400)>=600,color:col(st.color)||'#262626'};
  const s=n.data;
  for(let i=0;i<s.length;i++){
   rg.setStart(n,i);rg.setEnd(n,i+1);const rs=rg.getClientRects();if(!rs.length)continue;const r=rs[0];
   if(r.width<=0.01&&/\s/.test(s[i]))continue;
   const ch=/\s/.test(s[i])?' ':s[i];
   const top=r.top-F.top;
   if(!cur||forced||top>cur.ctop+cur.ch*0.5){cur={ctop:top,ch:r.height,chars:[]};lines.push(cur);forced=false;}
   cur.chars.push({c:ch,l:r.left-F.left,r:r.right-F.left,t:top,b:r.bottom-F.top,sty:sty});
  }
 }
 const L=[];
 for(const ln of lines){
  while(ln.chars.length&&ln.chars[0].c===' ')ln.chars.shift();
  while(ln.chars.length&&ln.chars[ln.chars.length-1].c===' ')ln.chars.pop();
  if(!ln.chars.length)continue;
  const runs=[];for(const c of ln.chars){const k=JSON.stringify(c.sty);
   if(runs.length&&runs[runs.length-1].k===k)runs[runs.length-1].t+=c.c;else runs.push({k:k,t:c.c,size:c.sty.size,bold:c.sty.bold,color:c.sty.color});}
  runs.forEach(r=>delete r.k);
  const xs=ln.chars.map(c=>c.l),xe=ln.chars.map(c=>c.r),ys=ln.chars.map(c=>c.t),ye=ln.chars.map(c=>c.b);
  L.push({x:Math.min(...xs),r:Math.max(...xe),y:Math.min(...ys),b:Math.max(...ye),size:Math.max(...runs.map(r=>r.size)),runs:runs});
 }
 if(!L.length)return;
 let al=cs.textAlign;al=(al==='center')?'center':((al==='right'||al==='end')?'right':'left');
 const lh=parseFloat(cs.lineHeight);
 const pitch=L.length>1?(L[L.length-1].y-L[0].y)/(L.length-1):(isNaN(lh)?L[0].size/0.75*1.4:lh);
 out.items.push({type:'text',align:al,pitch:pitch,lines:L});
 for(const ln of L){for(const r of ln.runs){if(r.size<11.9)out.warn.push('12pt 未満の文字: '+r.t.slice(0,16)+' ('+r.size.toFixed(1)+'pt)');}
  if(ln.r>F.width+2||ln.b>F.height+2||ln.x<-2||ln.y<-2)out.warn.push('図の外にはみ出す文字: '+ln.runs.map(r=>r.t).join('').slice(0,16));}
 if(el.scrollWidth>el.clientWidth+2&&el.clientWidth>0&&cs.overflow==='visible'&&el!==fig)
  out.warn.push('箱から文字がはみ出す: '+(el.textContent||'').trim().slice(0,16));
}
function walk(el){
 const cs=getComputedStyle(el);
 if(cs.display==='none'||cs.visibility==='hidden')return;
 if(el.tagName==='X-WIRE'||el.tagName==='SCRIPT'||el.tagName==='STYLE')return;
 if(el!==fig)emitBox(el,cs);
 if(el.classList&&el.classList.contains('ico')){const r=rel(el.getBoundingClientRect());
  out.items.push({type:'icon',kind:el.getAttribute('data-icon'),x:r.x,y:r.y,w:r.w,h:r.h,color:el.getAttribute('data-color')||null});return;}
 emitText(el,cs);
 for(const ch of el.children){if(!inl(ch))walk(ch);}
}
walk(fig);
for(const it of out.items){if(it.type==='rect'&&(it.x+it.w>F.width+2||it.y+it.h>F.height+2))out.warn.push('図の外にはみ出す箱');}
document.getElementById('__layout').textContent=JSON.stringify(out);
if(window.__PREVIEW){ // 保存した .html をブラウザで開いたとき、矢印も描いて見せる
 const NS='http://www.w3.org/2000/svg';const sv=document.createElementNS(NS,'svg');
 sv.setAttribute('width',F.width);sv.setAttribute('height',F.height);
 sv.style.cssText='position:absolute;left:0;top:0;overflow:visible;pointer-events:none';
 sv.innerHTML='<defs><marker id="ah" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="5" markerHeight="5" orient="auto-start-reverse"><path d="M0,0L10,5L0,10z" fill="context-stroke"/></marker></defs>';
 for(const w of out.wires){const pl=document.createElementNS(NS,'polyline');
  pl.setAttribute('points',w.pts.map(p=>p.join(',')).join(' '));pl.setAttribute('fill','none');
  pl.setAttribute('stroke',w.color);pl.setAttribute('stroke-width',w.width/0.75);
  if(w.dash)pl.setAttribute('stroke-dasharray','6 4');
  if(w.head==='end'||w.head==='both')pl.setAttribute('marker-end','url(#ah)');
  if(w.head==='start'||w.head==='both')pl.setAttribute('marker-start','url(#ah)');sv.appendChild(pl);}
 fig.appendChild(sv);
}
})();
"""


def _page(body, css, w_in, h_in):
    return ("<!doctype html><html><head><meta charset='utf-8'><style>" + BASE_CSS + "</style>"
            "<style>#fig{width:%.4fin;height:%.4fin}</style><style>%s</style></head><body>"
            "<div id='fig'>%s</div><pre id='__layout'></pre><script>%s</script></body></html>"
            % (w_in, h_in, css or "", body, _MEASURE_JS))


def measure(body, css, w_in, h_in, keep=None):
    """HTML を描画して配置を測る。keep を渡すと HTML(編集可能なマスター)をそこに保存する。"""
    page = _page(body, css, w_in, h_in)
    exe = next((p for p in _candidates() if p and os.path.exists(p)), None)
    if exe is None:
        raise SystemExit("ヘッドレス描画に使える Edge / Chrome が見つからない")
    tmp = tempfile.mkdtemp(prefix="corp_html_shapes_")
    try:
        hp = Path(tmp) / "fig.html"
        hp.write_text(page, encoding="utf-8")
        if keep:   # 編集用のマスター。ブラウザで開くと矢印まで描いた見た目になる
            Path(keep).parent.mkdir(parents=True, exist_ok=True)
            Path(keep).write_text(page.replace("<pre id='__layout'></pre><script>",
                                               "<pre id='__layout' style='display:none'></pre>"
                                               "<script>window.__PREVIEW=true;"), encoding="utf-8")
        wpx, hpx = int(w_in * PX) + 80, int(h_in * PX) + 80
        r = subprocess.run([exe, "--headless", "--disable-gpu", "--hide-scrollbars", "--dump-dom",
                            "--virtual-time-budget=3000", f"--window-size={wpx},{hpx}",
                            f"--user-data-dir={os.path.join(tmp, 'prof')}", hp.as_uri()],
                           capture_output=True, timeout=120)
        dom = (r.stdout or b"").decode("utf-8", "replace")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    i = dom.find('<pre id="__layout">')
    j = dom.find("</pre>", i)
    if i < 0 or j < 0:
        raise SystemExit("HTML の配置を測れなかった(ブラウザの出力に結果が無い)\n" + dom[-400:])
    raw = _html.unescape(dom[i + len('<pre id="__layout">'):j])
    if not raw.strip():
        raise SystemExit("HTML の配置が空(スクリプトのエラーか #fig が無い)")
    return json.loads(raw)


# ------------------------------------------------------------------ 図形を置く
def _ln_end(shape, head):
    ln = shape.line._get_or_add_ln()
    for tag, on in (("a:headEnd", head in ("start", "both")), ("a:tailEnd", head in ("end", "both"))):
        if on:
            el = ln.makeelement(qn(tag), {"type": "triangle", "w": "lg", "len": "lg"})   # 概念図の ▶ に近い大きさ
            ln.append(el)


def _style_line(shape, color, wpt, dash):
    shape.line.color.rgb = C._rgb(color)
    shape.line.width = Pt(wpt)
    if dash:
        shape.line.dash_style = MSO_LINE.DASH


def _line(s, x1, y1, x2, y2, color, wpt, dash=False, head="none"):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _style_line(cn, color, wpt, dash)
    _ln_end(cn, head)
    return cn


def _polyline(s, pts, color, wpt, dash=False, head="end"):
    if len(pts) == 2:
        return _line(s, pts[0][0], pts[0][1], pts[1][0], pts[1][1], color, wpt, dash, head)
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    sp = fb.convert_to_shape()
    sp.fill.background()
    sp.shadow.inherit = False
    _style_line(sp, color, wpt, dash)
    _ln_end(sp, head)
    return sp


def _text(d, s, it, ox, oy, sc):
    lines = it["lines"]
    x0 = min(l["x"] for l in lines)
    x1 = max(l["r"] for l in lines)
    w = (x1 - x0) / PX * sc
    slack = max(0.10, w * 0.08)                 # 游ゴシックの字幅差で折り返さないための余裕
    al = it["align"]
    x = ox + x0 / PX * sc - {"left": 0.0, "center": slack / 2, "right": slack}[al]
    pitch_in = it["pitch"] / PX * sc
    # 行の中心(ブラウザの字面の中央)に PowerPoint の 1 行の中心を合わせる。1 行の箱 = ピッチ
    c0 = (lines[0]["y"] + lines[0]["b"]) / 2 / PX * sc
    y = oy + c0 - pitch_in / 2
    h = pitch_in * len(lines)
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w + slack), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[al]
    if al == "left" and len(lines) > 1:
        # 行頭の位置が 1 行目と 2 行目以降で違う(ぶら下げ・字下げ)ときは段落のインデントで再現する
        first_x = lines[0]["x"]
        rest_x = min(l["x"] for l in lines[1:])
        if abs(rest_x - first_x) > 2:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches((rest_x - x0) / PX * sc))))
            pPr.set("indent", str(int(Inches((first_x - rest_x) / PX * sc))))
    p.line_spacing = Pt(it["pitch"] * PT_PER_PX * sc)
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    for k, ln in enumerate(lines):
        if k > 0:
            br = p._p.add_br()
            rpr = br.get_or_add_rPr()
            rpr.set("sz", str(int(round(ln["size"] * sc, 1) * 100)))
            rpr.set("lang", "ja-JP")
        for r in ln["runs"]:
            run = p.add_run()
            run.text = r["t"]
            name = C.YU
            if not r["bold"] and getattr(d, "body_font", None):
                name = d.body_font
            C._set_font(run, round(r["size"] * sc * 2) / 2, r["bold"], r["color"] or C.INK, name)
    return tb


def place_html(d, s, body, x, y, w, h, css="", keep=None):
    """本文の (x, y, w, h) inch に HTML の図を置く。戻り値は測った配置(JSON)。"""
    lay = measure(body, css, w, h, keep=keep)
    sc = w / (lay["w"] / PX) if lay["w"] else 1.0
    for it in lay["items"]:
        t = it["type"]
        if t == "rect":
            X, Y, Wd, Ht = x + it["x"] / PX * sc, y + it["y"] / PX * sc, it["w"] / PX * sc, it["h"] / PX * sc
            if it.get("ellipse"):
                shp, adj = MSO_SHAPE.OVAL, None
            elif it.get("radius", 0) > 0.5:
                shp, adj = MSO_SHAPE.ROUNDED_RECTANGLE, min(0.5, it["radius"] / max(min(it["w"], it["h"]), 1))
            else:
                shp, adj = MSO_SHAPE.RECTANGLE, None
            line = it.get("line")
            lw = it.get("lw", 1.0) * PT_PER_PX * sc if line else 1.0
            d.rect(s, X, Y, Wd, Ht, it.get("fill"), line=line, lw=lw, shape=shp, adj=adj, dash=it.get("dash", False))
        elif t == "line":
            _line(s, x + it["x1"] / PX * sc, y + it["y1"] / PX * sc, x + it["x2"] / PX * sc, y + it["y2"] / PX * sc,
                  it["color"], it["lw"] * PT_PER_PX * sc, it.get("dash", False))
        elif t == "text":
            _text(d, s, it, x, y, sc)
        elif t == "icon":
            size = min(it["w"], it["h"]) / PX * sc
            cx = x + (it["x"] + it["w"] / 2) / PX * sc
            cy = y + (it["y"] + it["h"] / 2) / PX * sc
            d.icon(s, it["kind"], cx - size / 2, cy - size / 2, size, it.get("color") or C.BLUE)
    for wv in lay["wires"]:
        pts = [(x + px_ / PX * sc, y + py_ / PX * sc) for px_, py_ in wv["pts"]]
        _polyline(s, pts, wv["color"], wv["width"], wv.get("dash", False), wv.get("head", "end"))
    for msg in lay.get("warn", []):
        d._warn("HTML 図: " + msg)
    return lay


def html_diagram(d, title, body, css="", lead=None, source=None, note=None, takeaway=None, keep=None):
    """HTML で組んだ図を 1 枚の本文ページにする(タイトル → ▍lead → 図 →(結論帯)→ 注・出典)。
    図のキャンバスは本文の全幅 × lead の下から本文下端まで。keep に .html を残すと後で直せる。"""
    s = d.slide(title)
    y0 = d.lead(s, lead) if lead else d.BODY_TOP + 0.15
    bot = d.body_bot(source, note, takeaway=bool(takeaway))
    top = y0 + 0.05
    place_html(d, s, body, d.MX, top, d.CW, bot - top - 0.08, css=css, keep=keep)
    if takeaway:
        d.takeaway_band(s, takeaway, source, note)
    d.source(s, source, note)
    return s
